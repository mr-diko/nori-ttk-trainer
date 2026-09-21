#!/usr/bin/env python3
"""
parse_ttk.py - Extracts dishes, preps, sets, decor tags, and embedded photos from 'ТТК NORI.xlsx'
and saves them into 'src/data/nori-menu.json' and 'public/images/'.
"""

import os
import re
import io
import json
import openpyxl
from PIL import Image

EXCEL_PATH = '/Users/admin/Downloads/ТТК NORI.xlsx'
OUTPUT_PATH = '/Users/admin/Downloads/nori-ttk-trainer/src/data/nori-menu.json'
PUBLIC_DIR = '/Users/admin/Downloads/nori-ttk-trainer/public'
DISH_IMG_DIR = os.path.join(PUBLIC_DIR, 'images', 'dishes')
SET_IMG_DIR = os.path.join(PUBLIC_DIR, 'images', 'sets')

DECOR_COLOR_HEX = 'FFE7F9EF'
DECOR_COLOR_ALT = 'FFAFE9CA'

# Known decor words to ensure robust tagging even if user edits colors
DECOR_KEYWORDS = [
    'мікрогрін', 'кунжут', 'нитки чилі', 'мигдальні пластівці', 
    'стружка тунця', 'боніто', 'ікра масаго', 'ікра імітована'
]

MAX_WIDTH = 800
MAX_HEIGHT = 800
JPEG_QUALITY = 82

def optimize_and_save(raw_bytes, target_path):
    os.makedirs(os.path.dirname(target_path), exist_ok=True)
    try:
        im = Image.open(io.BytesIO(raw_bytes))
        if im.mode in ('RGBA', 'LA', 'P'):
            bg = Image.new('RGB', im.size, (255, 255, 255))
            if im.mode == 'P':
                im = im.convert('RGBA')
            bg.paste(im, mask=im.split()[-1] if im.mode in ('RGBA', 'LA') else None)
            im = bg
        elif im.mode != 'RGB':
            im = im.convert('RGB')

        im.thumbnail((MAX_WIDTH, MAX_HEIGHT), Image.Resampling.LANCZOS)
        im.save(target_path, format='JPEG', quality=JPEG_QUALITY, optimize=True)
        return True
    except Exception as e:
        print(f"Error saving {target_path}: {e}")
        return False

def clean_str(val):
    if val is None:
        return ''
    s = str(val).strip()
    s = re.sub(r'[ \t]+', ' ', s)
    return s

def is_decor_cell(cell, ing_name):
    # Check cell fill
    if cell.fill and cell.fill.start_color and cell.fill.start_color.rgb:
        color = str(cell.fill.start_color.rgb).upper()
        if color in [DECOR_COLOR_HEX, DECOR_COLOR_ALT]:
            return True
    
    # Fallback to keyword heuristics for topping/garnish
    lower = ing_name.lower()
    for kw in DECOR_KEYWORDS:
        if kw in lower:
            return True
            
    return False

def parse_ttk():
    print(f"Loading workbook: {EXCEL_PATH} ...")
    wb = openpyxl.load_workbook(EXCEL_PATH)
    
    dishes = []
    preps = []
    sets = []
    
    dish_sheets = [
        'Філадельфія', 'Уромаки', 'Фірмові роли', 'Каліфорнія', 
        'Футохосомакі', 'Суші-бургери', 'Теплі, запечені, чіз роли', 
        'Рол доги', 'Норі роли', 'Оніґірі', 'Суші та гункани', 'Вегетаріанське АРХІВ'
    ]
    
    # 1. Parse Preps ('Заготовки')
    if 'Заготовки' in wb.sheetnames:
        sheet = wb['Заготовки']
        rows = list(sheet.iter_rows())
        current_prep = None
        prep_id_counter = 1
        
        for r in rows:
            c0 = r[0] if len(r) > 0 else None
            c1 = r[1] if len(r) > 1 else None
            c2 = r[2] if len(r) > 2 else None
            
            val0 = clean_str(c0.value) if c0 else ''
            val1 = clean_str(c1.value) if c1 else ''
            val2 = clean_str(c2.value) if c2 else ''
            
            if val1 == 'Технологія приготування' or (val0 and not val1 and not val2 and len(val0) < 50):
                if current_prep:
                    preps.append(current_prep)
                current_prep = {
                    'id': f'prep-{prep_id_counter}',
                    'name': val0,
                    'category': 'Заготовки',
                    'techProcess': '',
                    'ingredients': [],
                    'outputWeight': ''
                }
                prep_id_counter += 1
            elif current_prep:
                if 'Підготовка' in val0 or 'Зберігання' in val0 or len(val0) > 40:
                    current_prep['techProcess'] += (val0 + '\n')
                elif val0 == 'Вага готового продукту' and val1:
                    current_prep['outputWeight'] = val1
                elif val0 and val1:
                    current_prep['ingredients'].append({
                        'name': val0,
                        'weight': val1
                    })
        if current_prep:
            preps.append(current_prep)
            
    # 2. Parse Sets ('Набори 2026')
    if 'Набори 2026' in wb.sheetnames:
        sheet = wb['Набори 2026']
        images = getattr(sheet, '_images', [])
        rows = list(sheet.iter_rows())
        current_set = None
        current_set_row = -1
        set_id_counter = 1
        
        for idx, r in enumerate(rows):
            c0 = r[0] if len(r) > 0 else None
            c1 = r[1] if len(r) > 1 else None
            val0 = clean_str(c0.value) if c0 else ''
            val1 = clean_str(c1.value) if c1 else ''
            
            if val0:
                if current_set:
                    # Save image for previous set
                    matched = [img for img in images if hasattr(img.anchor, '_from') and abs(img.anchor._from.row - (current_set_row - 1)) <= 1]
                    if matched:
                        filename = f"{current_set['id']}.jpg"
                        target_path = os.path.join(SET_IMG_DIR, filename)
                        if optimize_and_save(matched[0]._data(), target_path):
                            current_set['image'] = f"images/sets/{filename}"
                    sets.append(current_set)
                
                weight_match = re.search(r'(\d+[\s]*г)', val0)
                pieces_match = re.search(r'\[(.*?)\]', val0)
                
                clean_name = re.sub(r'[\n\r]+', ' ', val0)
                clean_name = re.sub(r'\s+', ' ', clean_name).strip()
                
                current_set = {
                    'id': f'set-{set_id_counter}',
                    'name': clean_name,
                    'category': 'Набори 2026',
                    'totalWeight': weight_match.group(1) if weight_match else '',
                    'pieces': pieces_match.group(1) if pieces_match else '',
                    'rolls': []
                }
                current_set_row = idx
                set_id_counter += 1
            if val1 and current_set:
                current_set['rolls'].append(val1)

        if current_set:
            matched = [img for img in images if hasattr(img.anchor, '_from') and abs(img.anchor._from.row - (current_set_row - 1)) <= 1]
            if matched:
                filename = f"{current_set['id']}.jpg"
                target_path = os.path.join(SET_IMG_DIR, filename)
                if optimize_and_save(matched[0]._data(), target_path):
                    current_set['image'] = f"images/sets/{filename}"
            sets.append(current_set)
            
    # 3. Parse Dishes
    dish_id_counter = 1
    for sname in dish_sheets:
        if sname not in wb.sheetnames:
            continue
        sheet = wb[sname]
        images = getattr(sheet, '_images', [])
        rows = list(sheet.iter_rows())
        current_dish = None
        current_dish_row = -1
        
        for idx, r in enumerate(rows):
            c0 = r[0] if len(r) > 0 else None
            c1 = r[1] if len(r) > 1 else None
            c2 = r[2] if len(r) > 2 else None
            
            val0 = clean_str(c0.value) if c0 else ''
            val1 = clean_str(c1.value) if c1 else ''
            val2 = clean_str(c2.value) if c2 else ''
            
            if val0:
                # Commit previous dish
                if current_dish and len(current_dish['ingredients']) > 0:
                    matched = [img for img in images if hasattr(img.anchor, '_from') and abs(img.anchor._from.row - (current_dish_row - 1)) <= 1]
                    if matched:
                        filename = f"{current_dish['id']}.jpg"
                        target_path = os.path.join(DISH_IMG_DIR, filename)
                        if optimize_and_save(matched[0]._data(), target_path):
                            current_dish['image'] = f"images/dishes/{filename}"
                    dishes.append(current_dish)
                
                weight_match = re.search(r'(\d+[\s]*г)', val0)
                pieces_match = re.search(r'(\d+[\s]*шт)', val0)
                
                clean_name = re.sub(r'[\n\r]+', ' ', val0)
                clean_name = re.sub(r'\s+', ' ', clean_name).strip()
                
                current_dish = {
                    'id': f'dish-{dish_id_counter}',
                    'name': clean_name,
                    'category': sname,
                    'totalWeight': weight_match.group(1) if weight_match else '',
                    'pieces': pieces_match.group(1) if pieces_match else '',
                    'ingredients': []
                }
                current_dish_row = idx
                dish_id_counter += 1
                
            if val1 and current_dish:
                if val1 in ['Технологія приготування', 'Паніровка для ролів', 'Паніровка для бургерів']:
                    continue
                # Determine decor
                decor = is_decor_cell(c1, val1)
                current_dish['ingredients'].append({
                    'name': val1,
                    'weight': val2 if val2 else '',
                    'isDecor': decor
                })
                
        if current_dish and len(current_dish['ingredients']) > 0:
            matched = [img for img in images if hasattr(img.anchor, '_from') and abs(img.anchor._from.row - (current_dish_row - 1)) <= 1]
            if matched:
                filename = f"{current_dish['id']}.jpg"
                target_path = os.path.join(DISH_IMG_DIR, filename)
                if optimize_and_save(matched[0]._data(), target_path):
                    current_dish['image'] = f"images/dishes/{filename}"
            dishes.append(current_dish)
            
    all_categories = dish_sheets + ['Заготовки', 'Набори 2026']
    
    result = {
        'categories': all_categories,
        'dishCount': len(dishes),
        'prepCount': len(preps),
        'setCount': len(sets),
        'dishes': dishes,
        'preps': preps,
        'sets': sets
    }
    
    os.makedirs(os.path.dirname(OUTPUT_PATH), exist_ok=True)
    with open(OUTPUT_PATH, 'w', encoding='utf-8') as f:
        json.dump(result, f, ensure_ascii=False, indent=2)
        
    print(f"Successfully generated {OUTPUT_PATH}")
    print(f"Total dishes: {len(dishes)}")
    print(f"Total preps: {len(preps)}")
    print(f"Total sets: {len(sets)}")
    decor_count = sum(1 for d in dishes for i in d['ingredients'] if i.get('isDecor'))
    print(f"Total decor tagged ingredients: {decor_count}")
    dishes_with_img = sum(1 for d in dishes if d.get('image'))
    sets_with_img = sum(1 for s in sets if s.get('image'))
    print(f"Dishes with photo: {dishes_with_img}/{len(dishes)}")
    print(f"Sets with photo: {sets_with_img}/{len(sets)}")
    return result

if __name__ == '__main__':
    parse_ttk()
